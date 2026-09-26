from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

import os
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "tinh-gia-ve-sinh-kinh.xlsx")

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
GREEN = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HDR = Font(name=FONT, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
YELLOW = PatternFill("solid", fgColor="FFFF00")
KEY_FILL = PatternFill("solid", fgColor="E2EFDA")
SEC_FILL = PatternFill("solid", fgColor="DDEBF7")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

VND = '#,##0;(#,##0);"-"'
DEC = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'

wb = Workbook()


def style_all(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != FONT:
                c.font = Font(name=FONT, bold=c.font.bold, color=c.font.color, size=c.font.size)


def header(ws, row, labels):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font = HDR
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


# ---------------------------------------------------------------- Hướng dẫn
ws0 = wb.active
ws0.title = "HuongDan"
ws0["A1"] = "BẢNG TÍNH GIÁ DỊCH VỤ LAU KÍNH HẰNG NGÀY – NAM ĐẢO PHÚ QUỐC"
ws0["A1"].font = TITLE
lines = [
    "Mục đích: tính giá vốn và giá bán cho 1 m² kính mỗi lần lau, và dự báo lãi/lỗ 12 tháng đầu.",
    "",
    "CÁCH DÙNG",
    "1. Sửa các ô CHỮ XANH DƯƠNG (nền vàng = giả định quan trọng nhất) ở sheet GiaDinh, DauTu và dòng 'Số cửa hàng' ở sheet DuBao12T.",
    "2. Mọi ô chữ đen là công thức – không gõ đè. Ô chữ xanh lá là số lấy từ sheet khác.",
    "3. Xem kết quả ở sheet TinhGia (giá/m²), DoNhay (giá thay đổi thế nào khi năng suất thay đổi), DuBao12T (lãi/lỗ, vốn cần chuẩn bị).",
    "",
    "QUY ƯỚC",
    "• 'm²' = 1 m² của MỘT mặt kính (mặt ngoài, phía đường) cho MỘT lần lau.",
    "• Gói cơ bản: lau mặt ngoài mỗi ngày, 30 lần/tháng, kính cao ≤ 3 m (đứng dưới đất hoặc thang ngắn + sào).",
    "• Tất cả số tiền tính bằng VNĐ. Số liệu là ƯỚC TÍNH ban đầu – cần cập nhật sau khi khảo sát thực tế và chạy thử.",
    "",
    "SHEET",
    "GiaDinh  – toàn bộ giả định đầu vào (lương, năng suất, chi phí chung, thuế, lợi nhuận mục tiêu).",
    "DauTu    – danh sách trang thiết bị & chi phí khởi nghiệp.",
    "TinhGia  – cộng dồn chi phí ra giá vốn/m², giá bán/m², bảng giá các gói.",
    "DoNhay   – độ nhạy của giá theo tốc độ lau, tỷ lệ lấp đầy, diện tích kính/cửa hàng.",
    "DuBao12T – dự báo 12 tháng: số công nhân cần, doanh thu, chi phí, dòng tiền, tháng hoàn vốn.",
]
for i, t in enumerate(lines, 3):
    ws0.cell(row=i, column=1, value=t).font = BOLD if t.isupper() and t else BLACK
ws0.column_dimensions["A"].width = 130

# ---------------------------------------------------------------- Giả định
ws = wb.create_sheet("GiaDinh")
ws["A1"] = "GIẢ ĐỊNH ĐẦU VÀO (sửa ô chữ xanh)"
ws["A1"].font = TITLE
header(ws, 3, ["Nhóm", "Hạng mục", "Giá trị", "Đơn vị", "Ghi chú / nguồn"])
A = {}  # name -> absolute ref like GiaDinh!$C$5
r = 4


def section(title):
    global r
    c = ws.cell(row=r, column=1, value=title)
    c.font = BOLD
    for col in range(1, 6):
        ws.cell(row=r, column=col).fill = SEC_FILL
    r += 1


def inp(key, label, value, unit, note, fmt=VND, key_assump=False, formula=False):
    global r
    ws.cell(row=r, column=2, value=label).font = BLACK
    c = ws.cell(row=r, column=3, value=value)
    c.font = BLACK if formula else BLUE
    c.number_format = fmt
    c.border = BOX
    if key_assump:
        c.fill = YELLOW
    ws.cell(row=r, column=4, value=unit).font = BLACK
    ws.cell(row=r, column=5, value=note).font = BLACK
    ws.cell(row=r, column=5).alignment = Alignment(wrap_text=True, vertical="top")
    A[key] = f"GiaDinh!$C${r}"
    A[key + "_local"] = f"$C${r}"
    r += 1


section("1. Lương & chi phí 1 công nhân / tháng")
inp("luong", "Lương cơ bản", 7_000_000, "đ/tháng",
    "Thị trường tạp vụ 5–8 tr/tháng (2026); Phú Quốc thiếu lao động, làm sáng sớm ngoài trời → chọn 7 tr. Nguồn: CareerViet/Vieclam24h.", key_assump=True)
inp("xang", "Phụ cấp xăng xe + điện thoại", 700_000, "đ/tháng", "CN dùng xe máy cá nhân chạy giữa các cửa hàng.")
inp("chuyencan", "Thưởng chuyên cần / chất lượng", 500_000, "đ/tháng", "Chỉ trả khi không bị khách phàn nàn và đủ ảnh check-in.")
inp("luongbh", "Lương làm căn cứ đóng BH", 4_730_000, "đ/tháng",
    "Bằng lương tối thiểu vùng II năm 2026 (NĐ 293/2025). Đặc khu Phú Quốc thuộc vùng II.")
inp("tyle_bh", "Tỷ lệ BHXH+BHYT+BHTN doanh nghiệp đóng", 0.215, "%", "17,5% BHXH + 3% BHYT + 1% BHTN.", fmt=PCT)
inp("bhtn_ld", "Bảo hiểm tai nạn 24/24 (mua ngoài)", 25_000, "đ/tháng", "≈ 300.000 đ/người/năm.")
inp("thang13", "Trích lương tháng 13 / nghỉ phép", f"={A['luong_local']}/12", "đ/tháng", "Lương cơ bản / 12.", formula=True)
inp("chiphi_cn", "TỔNG CHI PHÍ 1 CÔNG NHÂN",
    f"={A['luong_local']}+{A['xang_local']}+{A['chuyencan_local']}+{A['luongbh_local']}*{A['tyle_bh_local']}+{A['bhtn_ld_local']}+{A['thang13_local']}",
    "đ/tháng", "Công thức.", formula=True)
ws.cell(row=r - 1, column=2).font = BOLD

section("2. Thời gian & năng suất")
inp("ngay", "Số ngày làm / tháng (mỗi CN)", 26, "ngày", "Nghỉ 4 ngày/tháng. Cửa hàng cần lau 30 ngày → xếp ca xoay vòng.", fmt="0")
inp("gio", "Giờ làm hiệu quả / ngày", 7, "giờ", "Ca 8 tiếng trừ họp đầu ca, lấy nước, nghỉ.", fmt="0.0")
inp("tocdo", "Tốc độ lau (mặt ngoài, bẩn nhẹ hằng ngày)", 1.0, "phút/m²",
    "≈ 60 m²/giờ gồm lau khung, cạnh. Thợ chuyên nghiệp 0,5–0,7; thợ mới 1,2–1,5. CẦN ĐO THỰC TẾ KHI CHẠY THỬ.", fmt="0.00", key_assump=True)
inp("codinh", "Thời gian cố định mỗi lần ghé 1 cửa hàng", 10, "phút",
    "Di chuyển giữa 2 cửa hàng gần nhau (~5 phút) + bày/dọn dụng cụ, chụp ảnh (~5 phút).", fmt="0", key_assump=True)
inp("m2shop", "Diện tích kính bình quân / cửa hàng (1 mặt)", 20, "m²",
    "Mặt tiền 5 m × cao 3 m + cửa ra vào ≈ 15–25 m². Đo thực tế 30 cửa hàng khi khảo sát.", fmt="0", key_assump=True)
inp("luot", "Số lần lau / tháng (gói hằng ngày)", 30, "lần", "Cửa hàng du lịch mở cửa cả tuần.", fmt="0")
inp("lapday", "Tỷ lệ lấp đầy (thời gian thực sự có việc)", 0.80, "%",
    "Phần còn lại là thời gian chết (khách hẹn lệch giờ, mưa, chờ mở cửa). Giai đoạn đầu thường 60–70%.", fmt=PCT, key_assump=True)

section("3. Vật tư, dụng cụ")
inp("khauhao", "Số tháng khấu hao bộ dụng cụ", 12, "tháng", "Dụng cụ ngoài trời, gió muối → dùng ~1 năm.", fmt="0")
inp("tieuhao", "Vật tư tiêu hao / CN / tháng", 300_000, "đ/tháng", "Lưỡi gạt cao su (2 cái), áo bông, khăn microfiber thay mới.")
inp("hoachat", "Hóa chất + nước / m²", 50, "đ/m²",
    "Nước lau kính đậm đặc ~60.000 đ/lít pha 1:100, ~0,04 lít dung dịch/m² → ~25 đ/m²; cộng hao hụt → 50 đ.")

section("4. Chi phí chung (quản lý thay chủ) / tháng")
inp("gd1_nguong", "Chuyển sang giai đoạn 2 khi số CN ≥", 6, "người", "Từ quy mô này cần 1 giám sát kiêm bán hàng toàn thời gian.", fmt="0")
inp("gd1", "Chi phí chung GĐ1 (≤ 5 CN)", 7_800_000, "đ/tháng",
    "Phụ cấp tổ trưởng 2 tr + marketing 2 tr + điện thoại/Zalo 0,5 tr + kế toán 0,8 tr + chỗ để đồ 1,5 tr + quỹ sự cố (vỡ/trầy kính) 1 tr.", key_assump=True)
inp("gd2", "Chi phí chung GĐ2 (≥ ngưỡng)", 20_000_000, "đ/tháng",
    "Giám sát/sales 12 tr + marketing 3 tr + điện thoại 0,5 tr + kế toán 1,5 tr + kho 1,5 tr + quỹ sự cố 1,5 tr.")
inp("quymo", "Quy mô tham chiếu để tính giá", 4, "công nhân", "Giá được tính khi đội đã ổn định ở quy mô này.", fmt="0", key_assump=True)

section("5. Thuế & lợi nhuận")
inp("thue_gia", "Dự phòng thuế trong giá bán", 0.05, "% doanh thu",
    "Hộ kinh doanh dịch vụ: GTGT 5% + TNCN 2% trên phần doanh thu vượt 500 tr/năm (từ 2026). Bình quân thực tế ~3–5%.", fmt=PCT)
inp("thue_tyle", "Thuế suất hộ KD dịch vụ (GTGT + TNCN)", 0.07, "%", "Dùng trong DuBao12T. Kiểm tra lại với cơ quan thuế đặc khu.", fmt=PCT)
inp("thue_nguong", "Ngưỡng doanh thu không chịu thuế", 500_000_000, "đ/năm", "Áp dụng từ 01/01/2026 cho hộ kinh doanh.")
inp("loinhuan", "Biên lợi nhuận mục tiêu (trên giá bán)", 0.25, "%", "Lợi nhuận của chủ – vì chủ không trực tiếp làm.", fmt=PCT, key_assump=True)

section("6. Hệ số giá các gói khác (so với gói hằng ngày)")
inp("hs_cachngay", "Gói cách ngày (15 lần/tháng)", 1.25, "x", "Kính bẩn hơn + chi phí di chuyển chia cho ít lần hơn.", fmt="0.00")
inp("hs_tuan", "Gói hằng tuần (4 lần/tháng)", 1.8, "x", "Kính bẩn nhiều, lau lâu hơn.", fmt="0.00")
inp("hs_tong", "Vệ sinh tổng 1 lần (2 mặt, cạo sơn/keo, lau khung)", 2.3, "x", "Thị trường 5.000–7.000 đ/m² (TP lớn, 2025).", fmt="0.00")
inp("hs_cao", "Phụ thu kính cao > 3 m", 1.5, "x", "Cần thang cao/sào dài, 2 người.", fmt="0.00")
inp("toithieu_m2", "Diện tích tính tối thiểu / lần ghé", 10, "m²", "Cửa hàng nhỏ vẫn tốn thời gian di chuyển.", fmt="0")

ws.column_dimensions["A"].width = 6
ws.column_dimensions["B"].width = 48
ws.column_dimensions["C"].width = 16
ws.column_dimensions["D"].width = 12
ws.column_dimensions["E"].width = 90
ws.freeze_panes = "A4"

# ---------------------------------------------------------------- Đầu tư
wd = wb.create_sheet("DauTu")
wd["A1"] = "TRANG THIẾT BỊ & CHI PHÍ KHỞI NGHIỆP (giá tham khảo 2026, VNĐ)"
wd["A1"].font = TITLE
header(wd, 3, ["Hạng mục", "Số lượng", "Đơn giá", "Thành tiền", "Ghi chú"])
kit = [
    ("Gạt kính cán inox 35–45 cm (loại tốt: Unger/Ettore hoặc tương đương)", 2, 350_000, "1 cái dùng, 1 cái dự phòng"),
    ("Bông lau kính 35 cm kèm 2 áo bông thay", 1, 450_000, "Strip washer"),
    ("Sào nối dài 3 đoạn 1,2–3,6 m", 1, 450_000, "Lau kính cao không cần thang"),
    ("Dao cạo kính + 10 lưỡi thay", 1, 150_000, "Cạo vết sơn, băng keo, phân chim"),
    ("Xô chữ nhật 20 L", 1, 200_000, "Vừa bông lau 35 cm"),
    ("Khăn microfiber (bộ 10 cái)", 1, 150_000, "Lau khung, mép"),
    ("Túi đeo hông đựng dụng cụ", 1, 200_000, ""),
    ("Bình xịt 1 L", 2, 30_000, ""),
    ("Thang nhôm chữ A 5 bậc (~1,5 m)", 1, 900_000, "Gọn, chở được trên xe máy"),
    ("Đồng phục in logo 3 bộ + nón + găng tay + áo mưa", 1, 750_000, "Tạo hình ảnh chuyên nghiệp"),
]
r = 4
wd.cell(row=r, column=1, value="A. BỘ DỤNG CỤ CHO 1 CÔNG NHÂN").font = BOLD
r += 1
kit_start = r
for name, q, p, note in kit:
    wd.cell(row=r, column=1, value=name)
    wd.cell(row=r, column=2, value=q).font = BLUE
    wd.cell(row=r, column=3, value=p).font = BLUE
    wd.cell(row=r, column=4, value=f"=B{r}*C{r}")
    wd.cell(row=r, column=5, value=note)
    r += 1
kit_end = r - 1
wd.cell(row=r, column=1, value="Tổng 1 bộ dụng cụ / công nhân").font = BOLD
wd.cell(row=r, column=4, value=f"=SUM(D{kit_start}:D{kit_end})").font = BOLD
wd.cell(row=r, column=4).fill = KEY_FILL
KIT = f"DauTu!$D${r}"
kit_total_row = r
r += 2

wd.cell(row=r, column=1, value="B. KHỞI NGHIỆP").font = BOLD
r += 1
start_rows_begin = r
wd.cell(row=r, column=1, value="Số bộ dụng cụ mua ban đầu (= số CN tháng đầu)")
wd.cell(row=r, column=2, value=3).font = BLUE
wd.cell(row=r, column=2).fill = YELLOW
wd.cell(row=r, column=3, value=f"=D{kit_total_row}")
wd.cell(row=r, column=4, value=f"=B{r}*C{r}")
wd.cell(row=r, column=5, value="2 CN chính + 1 CN dự phòng/tổ trưởng")
KITS0 = f"DauTu!$B${r}"
r += 1
shared = [
    ("Thang nhôm rút gọn 3,8 m (dùng chung)", 1, 2_200_000, "Cho kính cao, vệ sinh tổng"),
    ("Sào dài 6 m (dùng chung)", 1, 1_200_000, ""),
    ("Hóa chất lau kính đậm đặc can 5 L", 2, 300_000, "Đủ dùng vài tháng"),
    ("Điện thoại hotline + SIM số đẹp", 1, 2_500_000, "Zalo công việc, nhận đơn"),
    ("Đăng ký hộ kinh doanh, con dấu, TK ngân hàng", 1, 500_000, "UBND đặc khu Phú Quốc (một cửa)"),
    ("Marketing khai trương (tờ rơi 1.000 tờ, danh thiếp, Google Maps, fanpage)", 1, 3_000_000, ""),
    ("Chạy thử miễn phí 7 ngày cho 15 cửa hàng (vật tư, quà)", 1, 2_000_000, "Nhân công đã tính trong lương"),
]
for name, q, p, note in shared:
    wd.cell(row=r, column=1, value=name)
    wd.cell(row=r, column=2, value=q).font = BLUE
    wd.cell(row=r, column=3, value=p).font = BLUE
    wd.cell(row=r, column=4, value=f"=B{r}*C{r}")
    wd.cell(row=r, column=5, value=note)
    r += 1
wd.cell(row=r, column=1, value="Dự phòng phát sinh")
wd.cell(row=r, column=2, value=0.10).font = BLUE
wd.cell(row=r, column=2).number_format = PCT
wd.cell(row=r, column=4, value=f"=B{r}*SUM(D{start_rows_begin}:D{r-1})")
r += 1
wd.cell(row=r, column=1, value="TỔNG ĐẦU TƯ BAN ĐẦU (chưa gồm vốn lưu động)").font = BOLD
wd.cell(row=r, column=4, value=f"=SUM(D{start_rows_begin}:D{r-1})").font = BOLD
wd.cell(row=r, column=4).fill = KEY_FILL
CAPEX0 = f"DauTu!$D${r}"
r += 2
wd.cell(row=r, column=1, value="Ghi chú: xe máy do công nhân tự túc (đã có phụ cấp xăng). Nếu mua xe cho CN: +15–20 tr/xe cũ.")
wd.cell(row=r + 1, column=1, value="Chưa đầu tư máy nước tinh khiết (pure water, 30–60 tr) – chỉ cân nhắc khi có hợp đồng kính cao tầng/resort.")
for row in wd.iter_rows(min_row=4, max_row=r, min_col=2, max_col=4):
    for c in row:
        if c.number_format == "General":
            c.number_format = VND
wd.column_dimensions["A"].width = 70
wd.column_dimensions["B"].width = 11
wd.column_dimensions["C"].width = 14
wd.column_dimensions["D"].width = 16
wd.column_dimensions["E"].width = 45

# ---------------------------------------------------------------- Tính giá
wt = wb.create_sheet("TinhGia")
wt["A1"] = "TÍNH GIÁ VỐN & GIÁ BÁN / m² / LẦN LAU (gói hằng ngày, mặt ngoài)"
wt["A1"].font = TITLE
header(wt, 3, ["Bước", "Chỉ tiêu", "Giá trị", "Đơn vị", "Cách tính"])
T = {}
r = 4


def trow(key, step, label, formula, unit, how, fmt=VND, bold=False, fill=None, font=None):
    global r
    wt.cell(row=r, column=1, value=step)
    wt.cell(row=r, column=2, value=label).font = BOLD if bold else BLACK
    c = wt.cell(row=r, column=3, value=formula)
    c.number_format = fmt
    c.border = BOX
    c.font = font or (Font(name=FONT, bold=True) if bold else BLACK)
    if fill:
        c.fill = fill
    wt.cell(row=r, column=4, value=unit)
    wt.cell(row=r, column=5, value=how)
    T[key] = f"$C${r}"
    T[key + "_x"] = f"TinhGia!$C${r}"
    r += 1


trow("cp_cn", "1", "Chi phí 1 công nhân / tháng", f"={A['chiphi_cn']}", "đ/tháng", "Từ GiaDinh", font=GREEN)
trow("phut", "2", "Phút làm việc hiệu quả / CN / tháng", f"={A['ngay']}*{A['gio']}*60", "phút", "ngày × giờ × 60", fmt="#,##0")
trow("dg_phut", "3", "Chi phí nhân công / phút có mặt", f"={T['cp_cn']}/{T['phut']}", "đ/phút", "(1) / (2)")
trow("dc_thang", "4", "Dụng cụ (khấu hao) + tiêu hao / CN / tháng", f"={KIT}/{A['khauhao']}+{A['tieuhao']}", "đ/tháng", "bộ dụng cụ / số tháng + tiêu hao")
trow("dc_phut", "5", "Dụng cụ / phút có mặt", f"={T['dc_thang']}/{T['phut']}", "đ/phút", "(4) / (2)", fmt=DEC)
trow("tg_m2", "6", "Thời gian thực tế cho 1 m² (gồm phần di chuyển)", f"={A['tocdo']}+{A['codinh']}/{A['m2shop']}", "phút/m²",
     "tốc độ lau + thời gian cố định / m² mỗi cửa hàng", fmt="0.00")
trow("nl_cn", "7", "Năng lực 1 CN (m² lau được / tháng)", f"={T['phut']}*{A['lapday']}/{T['tg_m2']}", "m²/tháng", "(2) × lấp đầy / (6)", fmt="#,##0")
trow("nl_ngay", "", "  … tương đương m² / ngày làm", f"={T['nl_cn']}/{A['ngay']}", "m²/ngày", "", fmt="#,##0")
trow("shop_ngay", "", "  … tương đương số cửa hàng / ngày làm", f"={T['nl_ngay']}/{A['m2shop']}", "cửa hàng", "", fmt="0.0")
trow("tong_m2", "8", "Tổng m²/tháng ở quy mô tham chiếu", f"={T['nl_cn']}*{A['quymo']}", "m²/tháng", "(7) × số CN tham chiếu", fmt="#,##0")
trow("cpc", "9", "Chi phí chung / tháng ở quy mô tham chiếu",
     f"=IF({A['quymo']}>={A['gd1_nguong']},{A['gd2']},{A['gd1']})", "đ/tháng", "GĐ1 hoặc GĐ2 theo số CN")
r += 1
wt.cell(row=r, column=2, value="GIÁ VỐN / m² / LẦN").font = BOLD
for col in range(1, 6):
    wt.cell(row=r, column=col).fill = SEC_FILL
r += 1
trow("v_nc", "10", "Nhân công", f"={T['dg_phut']}*{T['tg_m2']}/{A['lapday']}", "đ/m²", "(3) × (6) / lấp đầy")
trow("v_dc", "11", "Dụng cụ & vật tư tiêu hao", f"={T['dc_phut']}*{T['tg_m2']}/{A['lapday']}", "đ/m²", "(5) × (6) / lấp đầy")
trow("v_hc", "12", "Hóa chất + nước", f"={A['hoachat']}", "đ/m²", "Từ GiaDinh", font=GREEN)
trow("v_cc", "13", "Chi phí chung (quản lý, marketing, kho…)", f"={T['cpc']}/{T['tong_m2']}", "đ/m²", "(9) / (8)")
trow("giavon", "14", "GIÁ VỐN / m²", f"=SUM({T['v_nc']}:{T['v_cc']})", "đ/m²", "(10)+(11)+(12)+(13)", bold=True, fill=KEY_FILL)
r += 1
wt.cell(row=r, column=2, value="GIÁ BÁN").font = BOLD
for col in range(1, 6):
    wt.cell(row=r, column=col).fill = SEC_FILL
r += 1
trow("giaban_tho", "15", "Giá bán tối thiểu để đạt lợi nhuận mục tiêu",
     f"={T['giavon']}/(1-{A['thue_gia']}-{A['loinhuan']})", "đ/m²", "giá vốn / (1 − thuế − lợi nhuận)")
trow("giaban", "16", "GIÁ BÁN ĐỀ XUẤT (làm tròn lên 100 đ)", f"=ROUNDUP({T['giaban_tho']}/100,0)*100", "đ/m²/lần",
     "Gói hằng ngày, mặt ngoài, kính ≤ 3 m", bold=True, fill=YELLOW)
trow("hoavon", "17", "Giá hòa vốn (lợi nhuận = 0)", f"={T['giavon']}/(1-{A['thue_gia']})", "đ/m²", "Không bán thấp hơn giá này")
trow("ln_m2", "18", "Lợi nhuận / m² ở giá đề xuất", f"={T['giaban']}*(1-{A['thue_gia']})-{T['giavon']}", "đ/m²", "")
trow("ln_thang", "19", "Lợi nhuận / tháng ở quy mô tham chiếu", f"={T['ln_m2']}*{T['tong_m2']}", "đ/tháng", "(18) × (8)", bold=True)
trow("dt_thang", "20", "Doanh thu / tháng ở quy mô tham chiếu", f"={T['giaban']}*{T['tong_m2']}", "đ/tháng", "")
trow("so_shop", "21", "Số cửa hàng cần có (gói hằng ngày) để lấp đầy", f"={T['tong_m2']}/({A['m2shop']}*{A['luot']})", "cửa hàng", "", fmt="0")
r += 1
wt.cell(row=r, column=2, value="BẢNG GIÁ GỢI Ý (đ/m²/lần)").font = BOLD
for col in range(1, 6):
    wt.cell(row=r, column=col).fill = SEC_FILL
r += 1
trow("p_ngay", "", "Gói hằng ngày (30 lần/tháng)", f"={T['giaban']}", "đ/m²", "")
trow("p_cach", "", "Gói cách ngày (15 lần/tháng)", f"=ROUNDUP({T['giaban']}*{A['hs_cachngay']}/100,0)*100", "đ/m²", "× hệ số GiaDinh")
trow("p_tuan", "", "Gói hằng tuần (4 lần/tháng)", f"=ROUNDUP({T['giaban']}*{A['hs_tuan']}/100,0)*100", "đ/m²", "× hệ số GiaDinh")
trow("p_tong", "", "Vệ sinh tổng 1 lần (2 mặt)", f"=ROUNDUP({T['giaban']}*{A['hs_tong']}/500,0)*500", "đ/m²", "× hệ số, làm tròn 500 đ")
trow("p_cao", "", "Phụ thu kính cao > 3 m (nhân với giá gói)", f"={A['hs_cao']}", "x", "", fmt="0.00")
r += 1
wt.cell(row=r, column=2, value="VÍ DỤ: CỬA HÀNG CÓ DIỆN TÍCH KÍNH BÌNH QUÂN").font = BOLD
for col in range(1, 6):
    wt.cell(row=r, column=col).fill = SEC_FILL
r += 1
trow("vd_m2", "", "Diện tích kính mặt ngoài", f"={A['m2shop']}", "m²", "", fmt="0", font=GREEN)
trow("vd_lan", "", "Phí mỗi lần lau", f"=MAX({T['vd_m2']},{A['toithieu_m2']})*{T['giaban']}", "đ/lần", "")
trow("vd_thang", "", "Phí gói hằng ngày / tháng", f"={T['vd_lan']}*{A['luot']}", "đ/tháng", "", bold=True, fill=KEY_FILL)
trow("vd_cach", "", "Phí gói cách ngày / tháng", f"=MAX({T['vd_m2']},{A['toithieu_m2']})*{T['p_cach']}*15", "đ/tháng", "")
trow("vd_tuan", "", "Phí gói hằng tuần / tháng", f"=MAX({T['vd_m2']},{A['toithieu_m2']})*{T['p_tuan']}*4", "đ/tháng", "")
trow("vd_min", "", "Phí tối thiểu gói hằng ngày (cửa hàng nhỏ)", f"={A['toithieu_m2']}*{T['giaban']}*{A['luot']}", "đ/tháng", "")

wt.column_dimensions["A"].width = 6
wt.column_dimensions["B"].width = 52
wt.column_dimensions["C"].width = 16
wt.column_dimensions["D"].width = 12
wt.column_dimensions["E"].width = 45
wt.freeze_panes = "A4"

# ---------------------------------------------------------------- Độ nhạy
wn = wb.create_sheet("DoNhay")
wn["A1"] = "ĐỘ NHẠY: GIÁ BÁN ĐỀ XUẤT (đ/m²) THAY ĐỔI THẾ NÀO?"
wn["A1"].font = TITLE
wn["A3"] = "Bảng 1 – Giá bán theo TỐC ĐỘ LAU (dòng) và TỶ LỆ LẤP ĐẦY (cột). Các giả định khác lấy từ GiaDinh."
wn["A3"].font = BOLD
wn["A4"] = "phút/m² \\ lấp đầy"
wn["A4"].font = HDR
wn["A4"].fill = HDR_FILL
utils = [0.6, 0.7, 0.8, 0.9]
speeds = [0.6, 0.8, 1.0, 1.2, 1.5]
for j, u in enumerate(utils):
    c = wn.cell(row=4, column=2 + j, value=u)
    c.font = Font(name=FONT, bold=True, color="0000FF")
    c.number_format = PCT
    c.fill = SEC_FILL
for i, s in enumerate(speeds):
    rr = 5 + i
    c = wn.cell(row=rr, column=1, value=s)
    c.font = Font(name=FONT, bold=True, color="0000FF")
    c.number_format = "0.00"
    c.fill = SEC_FILL
    for j in range(len(utils)):
        col = get_column_letter(2 + j)
        time = f"($A{rr}+{A['codinh']}/{A['m2shop']})"
        u = f"{col}$4"
        cap = f"({T['phut_x']}*{u}/{time})"
        f = (f"=ROUNDUP((({T['dg_phut_x']}+{T['dc_phut_x']})*{time}/{u}+{A['hoachat']}+{T['cpc_x']}/({cap}*{A['quymo']}))"
             f"/(1-{A['thue_gia']}-{A['loinhuan']})/100,0)*100")
        cell = wn.cell(row=rr, column=2 + j, value=f)
        cell.number_format = VND
        cell.border = BOX

wn["A12"] = "Bảng 2 – Theo DIỆN TÍCH KÍNH / CỬA HÀNG (cửa hàng nhỏ tốn tiền di chuyển hơn trên mỗi m²)."
wn["A12"].font = BOLD
header(wn, 13, ["m² / cửa hàng", "Giá bán đ/m²", "Phí gói ngày / tháng", "Số cửa hàng / CN / ngày"])
for i, m in enumerate([8, 10, 15, 20, 30, 50]):
    rr = 14 + i
    c = wn.cell(row=rr, column=1, value=m)
    c.font = BLUE
    time = f"({A['tocdo']}+{A['codinh']}/A{rr})"
    cap = f"({T['phut_x']}*{A['lapday']}/{time})"
    wn.cell(row=rr, column=2, value=(f"=ROUNDUP((({T['dg_phut_x']}+{T['dc_phut_x']})*{time}/{A['lapday']}+{A['hoachat']}+{T['cpc_x']}/({cap}*{A['quymo']}))"
                                     f"/(1-{A['thue_gia']}-{A['loinhuan']})/100,0)*100")).number_format = VND
    wn.cell(row=rr, column=3, value=f"=B{rr}*A{rr}*{A['luot']}").number_format = VND
    wn.cell(row=rr, column=4, value=f"={cap}/{A['ngay']}/A{rr}").number_format = "0.0"
wn["A21"] = ("Đọc bảng: nếu thợ chỉ lau được 1,5 phút/m² và lấp đầy 60%, giá phải lên mức ô tương ứng. "
             "→ Đào tạo tốc độ + gom khách theo tuyến là 2 đòn bẩy quan trọng nhất.")
wn.column_dimensions["A"].width = 22
for col in "BCDE":
    wn.column_dimensions[col].width = 22

# ---------------------------------------------------------------- Dự báo 12 tháng
wf = wb.create_sheet("DuBao12T")
wf["A1"] = "DỰ BÁO 12 THÁNG ĐẦU (dòng tiền, VNĐ)"
wf["A1"].font = TITLE
wf["A2"] = "Sửa dòng 'Số cửa hàng ký gói ngày' (chữ xanh) theo kế hoạch bán hàng thực tế. Tháng 0 = chuẩn bị & mua sắm."
months = list(range(0, 13))
header(wf, 4, ["Chỉ tiêu"] + [f"T{m}" for m in months])
ncol = len(months)
rows = {}
r = 5
shops = [0, 6, 12, 18, 24, 30, 35, 40, 45, 50, 55, 60, 65]
extra = [0, 0, 1_000_000, 2_000_000, 3_000_000, 3_000_000, 4_000_000, 4_000_000, 5_000_000, 5_000_000, 6_000_000, 6_000_000, 6_000_000]


def frow(key, label, vals, fmt=VND, bold=False, font=None, fill=None):
    global r
    wf.cell(row=r, column=1, value=label).font = BOLD if bold else BLACK
    for j, v in enumerate(vals):
        c = wf.cell(row=r, column=2 + j, value=v)
        c.number_format = fmt
        c.font = font or (Font(name=FONT, bold=True) if bold else BLACK)
        if fill:
            c.fill = fill
    rows[key] = r
    r += 1


def col(j):
    return get_column_letter(2 + j)


frow("shop", "Số cửa hàng ký gói ngày (cuối tháng)", shops, fmt="0", font=BLUE, fill=YELLOW)
frow("m2", "m² lau / tháng", [f"={col(j)}{rows['shop']}*{A['m2shop']}*{A['luot']}" for j in months], fmt="#,##0")
frow("phut", "Phút lao động cần", [f"={col(j)}{rows['shop']}*{A['luot']}*({A['codinh']}+{A['m2shop']}*{A['tocdo']})" for j in months], fmt="#,##0")
frow("cn", "Số công nhân cần (tối thiểu 2 để xoay ca 30 ngày)",
     [0] + [f"=MAX(2,ROUNDUP({col(j)}{rows['phut']}/({T['phut_x']}*{A['lapday']}),0))" for j in months[1:]], fmt="0")
wf.cell(row=rows["cn"], column=2).font = BLUE
r += 1
frow("dt1", "Doanh thu gói hằng ngày", [f"={col(j)}{rows['m2']}*{T['giaban_x']}" for j in months])
frow("dt2", "Doanh thu vệ sinh tổng / dịch vụ thêm", extra, font=BLUE)
frow("dt", "TỔNG DOANH THU", [f"={col(j)}{rows['dt1']}+{col(j)}{rows['dt2']}" for j in months], bold=True)
r += 1
frow("c_nc", "Lương & bảo hiểm công nhân", [f"={col(j)}{rows['cn']}*{A['chiphi_cn']}" for j in months])
frow("c_vt", "Vật tư tiêu hao", [f"={col(j)}{rows['cn']}*{A['tieuhao']}" for j in months])
frow("c_hc", "Hóa chất + nước", [f"={col(j)}{rows['m2']}*{A['hoachat']}" for j in months])
frow("c_cc", "Chi phí chung (quản lý, marketing, kho…)",
     [0] + [f"=IF({col(j)}{rows['cn']}>={A['gd1_nguong']},{A['gd2']},{A['gd1']})" for j in months[1:]])
frow("c_thue", "Thuế hộ KD (trên phần DT vượt ngưỡng, ước theo tháng)",
     [f"=MAX(0,{col(j)}{rows['dt']}-{A['thue_nguong']}/12)*{A['thue_tyle']}" for j in months])
frow("cp", "TỔNG CHI PHÍ HOẠT ĐỘNG", [f"=SUM({col(j)}{rows['c_nc']}:{col(j)}{rows['c_thue']})" for j in months], bold=True)
frow("ln", "LỢI NHUẬN HOẠT ĐỘNG", [f"={col(j)}{rows['dt']}-{col(j)}{rows['cp']}" for j in months], bold=True, fill=KEY_FILL)
frow("bien", "Biên lợi nhuận", [f"=IF({col(j)}{rows['dt']}=0,0,{col(j)}{rows['ln']}/{col(j)}{rows['dt']})" for j in months], fmt=PCT)
r += 1
frow("capex", "Đầu tư (T0 = ban đầu; sau đó mua dụng cụ cho CN mới)",
     [f"={CAPEX0}"] + [f"=MAX(0,{col(j)}{rows['cn']}-{('%s' % KITS0) if j == 1 else col(j-1) + str(rows['cn'])})*{KIT}" for j in months[1:]])
frow("cf", "DÒNG TIỀN RÒNG", [f"={col(j)}{rows['ln']}-{col(j)}{rows['capex']}" for j in months], bold=True)
frow("cum", "Dòng tiền lũy kế", [f"={col(0)}{rows['cf']}"] + [f"={col(j-1)}{r}+{col(j)}{rows['cf']}" for j in months[1:]], bold=True)
frow("flag", "(phụ) tháng lũy kế chuyển dương",
     [""] + [f'=IF(AND({col(j)}{rows["cum"]}>=0,{col(j-1)}{rows["cum"]}<0),{j},"")' for j in months[1:]], fmt="0")
r += 1
last = col(12)
wf.cell(row=r, column=1, value="VỐN CẦN CHUẨN BỊ (đáy dòng tiền lũy kế)").font = BOLD
c = wf.cell(row=r, column=2, value=f"=-MIN(B{rows['cum']}:{last}{rows['cum']})")
c.number_format = VND
c.font = BOLD
c.fill = YELLOW
wf.cell(row=r, column=3, value="→ nên chuẩn bị thêm 20–30% dự phòng.")
rows["von"] = r
r += 1
wf.cell(row=r, column=1, value="Tháng hoàn vốn (lũy kế ≥ 0)").font = BOLD
c = wf.cell(row=r, column=2, value=f'=IF(MAX(C{rows["flag"]}:{last}{rows["flag"]})=0,"Chưa hoàn vốn",MIN(C{rows["flag"]}:{last}{rows["flag"]}))')
c.font = BOLD
c.fill = YELLOW
rows["hoan"] = r
r += 1
wf.cell(row=r, column=1, value="Lợi nhuận cả năm 1 (T1–T12)").font = BOLD
c = wf.cell(row=r, column=2, value=f"=SUM(C{rows['ln']}:{last}{rows['ln']})")
c.number_format = VND
c.font = BOLD
rows["ln_nam"] = r
r += 1
wf.cell(row=r, column=1, value="Doanh thu cả năm 1 (T1–T12)").font = BOLD
c = wf.cell(row=r, column=2, value=f"=SUM(C{rows['dt']}:{last}{rows['dt']})")
c.number_format = VND
c.font = BOLD
rows["dt_nam"] = r
wf.cell(row=rows["capex"], column=2).font = GREEN
wf.column_dimensions["A"].width = 55
for j in months:
    wf.column_dimensions[col(j)].width = 13
wf.freeze_panes = "B5"

for s in wb.worksheets:
    for row in s.iter_rows():
        for c in row:
            f = c.font
            if f.name != FONT:
                c.font = Font(name=FONT, bold=f.bold, color=f.color, size=f.size)

wb.save(OUT)
print("Đã tạo", OUT, "- mở bằng Excel/LibreOffice để tính lại công thức")
